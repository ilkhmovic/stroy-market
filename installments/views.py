import re
from datetime import datetime
from decimal import Decimal

from django.apps import apps
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import models as dj_models
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST

from . import services
from .models import NasiyaChoice, WalletCard


def _back(request, default='installments:cards'):
    """Amaldan keyin foydalanuvchini kelgan sahifasiga qaytaradi."""
    nxt = request.POST.get('next') or ''
    if nxt and url_has_allowed_host_and_scheme(nxt, allowed_hosts={request.get_host()},
                                               require_https=request.is_secure()):
        return redirect(nxt)
    return redirect(default)


# ───────────── MIJOZ: kartalar ─────────────

@login_required
def cards_page(request):
    Order = apps.get_model('main', 'Order')
    cards = request.user.wallet_cards.all()
    unpaid = []
    qs = services.wallet_paid_annotation(Order.objects.filter(user=request.user, status='pending'))
    for order in qs.filter(wallet_paid=False).order_by('-id')[:20]:
        total = services.order_total(order)
        if total > 0:
            unpaid.append({'order': order, 'total': total})
    return render(request, 'installments/cards.html', {'cards': cards, 'unpaid_orders': unpaid})


@login_required
@require_POST
def add_card(request):
    number = re.sub(r'\D', '', request.POST.get('card_number', ''))
    expiry = request.POST.get('expiry', '').strip()
    holder = re.sub(r'\s+', ' ', request.POST.get('holder_name', '')).strip()[:60]
    m = re.fullmatch(r'(\d{1,2})\s*/\s*(\d{2}|\d{4})', expiry)
    if len(number) != 16:
        messages.error(request, "Karta raqami 16 ta raqamdan iborat bo'lishi kerak.")
        return _back(request)
    if not m:
        messages.error(request, "Amal qilish muddatini OO/YY ko'rinishida kiriting (masalan 09/28).")
        return _back(request)
    month, year = int(m.group(1)), int(m.group(2))
    year = year + 2000 if year < 100 else year
    if not 1 <= month <= 12:
        messages.error(request, "Amal qilish oyi noto'g'ri.")
        return _back(request)
    card = WalletCard(user=request.user, holder_name=holder, brand=services.detect_brand(number),
                      last4=number[-4:], exp_month=month, exp_year=year)
    if card.is_expired:
        messages.error(request, "Kartaning amal qilish muddati tugagan.")
        return _back(request)
    if request.user.wallet_cards.filter(last4=card.last4, exp_month=month, exp_year=year,
                                        brand=card.brand).exists():
        messages.info(request, "Bu karta allaqachon qo'shilgan.")
        return _back(request)
    card.save()
    messages.success(request, "Karta qo'shildi. Hisobingizni menejer to'ldirgach, undagi pul bilan xarid qila olasiz.")
    return _back(request)


@login_required
def card_history(request, card_id):
    card = get_object_or_404(WalletCard, pk=card_id, user=request.user)
    return render(request, 'installments/card_history.html',
                  {'card': card, 'transactions': card.transactions.select_related('order')})


@login_required
@require_POST
def pay_order(request):
    try:
        order, total = services.pay_order(request.user, request.POST.get('card_id'), request.POST.get('order_id'))
        messages.success(request, f"Buyurtma #{order.pk} karta hisobidan to'landi: {services.fmt_sum(total)} so'm.")
    except services.CardError as e:
        messages.error(request, str(e))
    return _back(request)


# ───────────── SOTUVCHI: nasiya arizalari Excel ─────────────

def _date_lookup(model, candidates=('created_at', 'created', 'submitted_at', 'date')):
    for name in candidates:
        try:
            f = model._meta.get_field(name)
        except Exception:
            continue
        suffix = '__date' if isinstance(f, dj_models.DateTimeField) else ''
        return name, suffix
    return None, ''


def _parse_date(value):
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def _seller_applications(store, date_from, date_to):
    OrderItem = apps.get_model('main', 'OrderItem')
    App = apps.get_model('main', 'NasiyaApplication')
    order_ids = OrderItem.objects.filter(product__store=store).values('order_id')
    qs = App.objects.filter(order_id__in=order_ids).select_related('order', 'order__user')
    field, suffix = _date_lookup(App)
    prefix = ''
    if not field:
        field, suffix = _date_lookup(apps.get_model('main', 'Order'))
        prefix = 'order__'
    if field:
        if date_from:
            qs = qs.filter(**{f'{prefix}{field}{suffix}__gte': date_from})
        if date_to:
            qs = qs.filter(**{f'{prefix}{field}{suffix}__lte': date_to})
        qs = qs.order_by(f'-{prefix}{field}')
    return qs, (prefix + field if field else None)


def _seller_store(request):
    if not getattr(request.user, 'is_seller', False):
        return None
    return getattr(request.user, 'store', None)


@login_required
def seller_export(request):
    store = _seller_store(request)
    if store is None:
        messages.error(request, "Bu sahifa faqat do'kon egalari uchun.")
        return redirect('marketplace')
    date_from = _parse_date(request.GET.get('date_from'))
    date_to = _parse_date(request.GET.get('date_to'))
    qs, date_path = _seller_applications(store, date_from, date_to)
    if request.GET.get('download') == '1':
        return _build_xlsx(store, qs, date_path, date_from, date_to)
    return render(request, 'installments/seller_export.html', {
        'count': qs.count(),
        'date_from': request.GET.get('date_from', ''), 'date_to': request.GET.get('date_to', ''),
    })


def _dig(obj, path):
    for part in path.split('__'):
        obj = getattr(obj, part, None)
        if obj is None:
            return None
    return obj


def _build_xlsx(store, qs, date_path, date_from, date_to):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    OrderItem = apps.get_model('main', 'OrderItem')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Nasiya arizalari'
    headers = ['Ariza ID', 'Sana', 'Mijoz (F.I.Sh.)', 'Telefon', 'Ish joyi', 'Oylik daromad (so\'m)',
               'Nasiya hamkori', 'Muddat (oy)', 'Ustama (%)', 'Mahsulotlar', 'Mahsulotlar summasi (so\'m)',
               'Ustama bilan jami (so\'m)', 'Oyiga to\'lov (so\'m)', 'Ariza holati']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1D4ED8')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for app in qs:
        items = list(OrderItem.objects.filter(order=app.order, product__store=store))
        subtotal = sum((i.price * i.quantity for i in items), Decimal('0'))
        choice = NasiyaChoice.objects.filter(application=app).first()
        months = choice.months if choice else getattr(app, 'term_months', None)
        percent = choice.markup_percent if choice else None
        total = monthly = None
        if percent is not None and months:
            total, monthly = services.calc_installment(subtotal, percent, months)
        when = _dig(app, date_path) if date_path else None
        if when is not None and hasattr(when, 'tzinfo') and when.tzinfo:
            from django.utils import timezone
            when = timezone.localtime(when).replace(tzinfo=None)
        monthly_income = getattr(app, 'monthly_income', None)
        ws.append([
            app.pk, when,
            getattr(app, 'full_name', '') or app.order.user.get_full_name(),
            getattr(app, 'phone', ''), getattr(app, 'workplace', ''),
            float(monthly_income) if monthly_income is not None else None,
            choice.partner_name if choice else '', months,
            float(percent) if percent is not None else None,
            ', '.join(f"{i.product_name} x{i.quantity}" for i in items),
            float(subtotal),
            float(total) if total is not None else None,
            float(monthly) if monthly is not None else None,
            app.get_status_display() if hasattr(app, 'get_status_display') else '',
        ])
    if ws.max_row == 1:
        ws.append(["Tanlangan davrda ma'lumot yo'q"])

    for idx in range(1, len(headers) + 1):
        letter = get_column_letter(idx)
        width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 60)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = 'DD.MM.YYYY HH:MM'
    for col in ('F', 'K', 'L', 'M'):
        for c in ws[col][1:]:
            c.number_format = '#,##0'
    ws.freeze_panes = 'A2'

    period = f"{date_from or 'boshidan'}_{date_to or 'hozirgacha'}"
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="nasiya_arizalari_{period}.xlsx"'
    wb.save(resp)
    return resp
